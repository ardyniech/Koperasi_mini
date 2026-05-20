#!/usr/bin/env python3
from openpyxl import load_workbook

wb = load_workbook('/home/ardy/koperasi_mini/koperasi_FINAL_with_FORMULAS.xlsx')

print("=== CEK FORMULAS ===")
print("\n1. Dashboard Sheet:")
ws = wb['Dashboard']
for row in range(3, 7):
    cell = ws.cell(row=row, column=2)
    is_formula = str(cell.value).startswith('=') if cell.value else False
    print(f"  Row {row}: {cell.value} | Formula: {is_formula}")

print("\n2. DCF Sheet (NPV & IRR):")
ws = wb['DCF']
print(f"  NPV cell B9: {ws['B9'].value}")
print(f"  IRR cell B10: {ws['B10'].value}")

print("\n3. Amortization Sheet (PMT):")
ws = wb['Amortization']
print(f"  PMT cell B6: {ws['B6'].value}")

print("\n4. Members Sheet (IF & PD):")
ws = wb['Members']
for row in [4, 5, 6]:
    print(f"  Row {row}, Risk (I): {ws.cell(row=row, column=9).value}")
    print(f"  Row {row}, PD (J): {ws.cell(row=row, column=10).value}")

print("\n=== HASIL AKHIR ===")
count_formulas = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                count_formulas += 1

print(f"Total formulas found: {count_formulas}")
if count_formulas > 0:
    print("✅ EXCEL BENERAN ADA FORMULAS!")
    print("Formulas akan otomatis menghitung saat file dibuka di Excel!")
else:
    print("❌ GAK ADA FORMULAS - NGIBUL LAGI!")
